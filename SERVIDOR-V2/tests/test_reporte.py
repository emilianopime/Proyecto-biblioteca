"""
Pruebas del reporte de estadisticas: rango de periodos (sin base de datos),
calculo, y las tres salidas (pagina, PDF, Excel) contra un Postgres real.
"""
import io
import os
from datetime import date, datetime

import pytest

from reporte import rango_periodo, PeriodoInvalido

# ---------------------------------------------------------------------------
# Periodos (puro)
# ---------------------------------------------------------------------------

def test_este_mes_va_del_primero_al_dia_de_hoy():
    desde, hasta, etiqueta = rango_periodo("mes", hoy=date(2026, 9, 18))
    assert (desde, hasta) == (date(2026, 9, 1), date(2026, 9, 18))
    assert etiqueta == "Septiembre de 2026"


def test_mes_anterior_es_el_mes_completo():
    desde, hasta, etiqueta = rango_periodo("mes-anterior", hoy=date(2026, 1, 10))
    assert (desde, hasta) == (date(2025, 12, 1), date(2025, 12, 31))
    assert etiqueta == "Diciembre de 2025"


@pytest.mark.parametrize("hoy, inicio, etiqueta", [
    (date(2026, 9, 18), date(2026, 8, 1), "Semestre agosto-diciembre 2026"),
    (date(2026, 3, 2),  date(2026, 1, 1), "Semestre enero-julio 2026"),
    (date(2026, 7, 15), date(2026, 1, 1), "Semestre enero-julio 2026"),
])
def test_semestre_es_enero_julio_o_agosto_diciembre(hoy, inicio, etiqueta):
    desde, hasta, et = rango_periodo("semestre", hoy=hoy)
    assert desde == inicio
    assert hasta == hoy
    assert et == etiqueta


def test_rango_explicito_conserva_las_fechas_y_las_etiqueta():
    desde, hasta, etiqueta = rango_periodo("rango", hoy=date(2026, 9, 18), desde="2026-08-15", hasta="2026-09-01")
    assert (desde, hasta) == (date(2026, 8, 15), date(2026, 9, 1))
    assert etiqueta == "Del 15/08/2026 al 01/09/2026"


def test_rango_invertido_o_mal_escrito_da_error_claro():
    with pytest.raises(PeriodoInvalido):
        rango_periodo("rango", hoy=date(2026, 9, 18), desde="2026-09-10", hasta="2026-09-01")
    with pytest.raises(PeriodoInvalido):
        rango_periodo("rango", hoy=date(2026, 9, 18), desde="ayer", hasta="2026-09-01")
    with pytest.raises(PeriodoInvalido):
        rango_periodo("trimestre", hoy=date(2026, 9, 18))


# ---------------------------------------------------------------------------
# Calculo y salidas (Postgres real)
# ---------------------------------------------------------------------------

DSN = os.getenv("TEST_DB_DSN")
con_db = pytest.mark.skipif(not DSN, reason="TEST_DB_DSN no definida")


@pytest.fixture(scope="module")
def app():
    if not DSN:
        pytest.skip("TEST_DB_DSN no definida")
    partes = dict(p.split("=", 1) for p in DSN.split())
    os.environ.update({"DB_HOST": partes["host"], "DB_PORT": partes["port"], "DB_USER": partes["user"],
                       "DB_PASSWORD": partes["password"], "DB_NAME": partes["dbname"]})
    import main
    main.DB_CONFIG = {"host": partes["host"], "port": partes["port"], "user": partes["user"],
                      "password": partes["password"], "database": partes["dbname"]}
    main.app.config["TESTING"] = True
    return main.app


@pytest.fixture
def datos(app):
    import psycopg2
    with psycopg2.connect(DSN) as conexion, conexion.cursor() as cur:
        cur.execute("DROP TABLE IF EXISTS bitacora_uso; DROP TABLE IF EXISTS alumnos;")
        cur.execute("CREATE TABLE alumnos (cardnumber BIGINT PRIMARY KEY, surname TEXT, firstname TEXT, sort1 TEXT);")
        cur.execute("CREATE TABLE bitacora_uso (id SERIAL PRIMARY KEY, computer_id TEXT, matricula BIGINT, evento TEXT, timestamp TIMESTAMP DEFAULT NOW());")
        cur.executemany("INSERT INTO alumnos VALUES (%s, %s, %s, %s);", [
            (1, "Perez", "Juan", "Derecho"), (2, "Lopez", "Ana", "Derecho"), (3, "Ruiz", "Luis", "Medicina")])
        # Las marcas se guardan en UTC; Chihuahua es UTC-6. 15:00 UTC = 09:00 local.
        eventos = [
            ("PC-01", 1, "LOGIN",  "2026-09-01 15:00:00"),
            ("PC-01", 1, "LOGOUT", "2026-09-01 16:00:00"),
            ("PC-02", 2, "LOGIN",  "2026-09-01 15:30:00"),
            ("PC-01", 3, "LOGIN",  "2026-09-02 20:00:00"),   # 14:00 local
            ("PC-03", 1, "LOGIN",  "2026-08-20 15:00:00"),   # fuera del periodo
            ("PC-01", 2, "LOGIN",  "2026-09-03 03:30:00"),   # 21:30 local del dia 2
        ]
        cur.executemany("INSERT INTO bitacora_uso (computer_id, matricula, evento, timestamp) VALUES (%s, %s, %s, %s);", eventos)
    with app.test_client() as c:
        with c.session_transaction() as s:
            s["authenticated"] = True
        yield c


@con_db
def test_calculo_del_reporte_en_un_rango(datos):
    d = datos.get("/api/reporte?periodo=rango&desde=2026-09-01&hasta=2026-09-30").get_json()

    assert d["periodo"]["etiqueta"] == "Del 01/09/2026 al 30/09/2026"
    assert d["total_entradas"] == 4          # el LOGIN de agosto no cuenta, el LOGOUT tampoco
    assert d["alumnos_distintos"] == 3
    assert d["dias_con_actividad"] == 2
    assert d["equipos"][0] == {"pc": "PC-01", "total": 3}
    assert d["carreras"][0] == {"carrera": "Derecho", "total": 3}
    assert d["horas"][9] == 2                # 09:00 local
    assert d["horas"][14] == 1 and d["horas"][21] == 1
    assert d["dias"] == [{"fecha": "2026-09-01", "total": 2}, {"fecha": "2026-09-02", "total": 2}]


@con_db
def test_periodo_invalido_devuelve_400(datos):
    r = datos.get("/api/reporte?periodo=rango&desde=2026-09-10&hasta=2026-09-01")
    assert r.status_code == 400
    assert "fecha" in r.get_json()["error"].lower()


@con_db
def test_pagina_del_reporte_es_html_con_el_periodo_y_las_cifras(datos):
    r = datos.get("/reporte?periodo=rango&desde=2026-09-01&hasta=2026-09-30")
    html = r.get_data(as_text=True)
    assert r.status_code == 200 and "text/html" in r.content_type
    assert "Del 01/09/2026 al 30/09/2026" in html
    assert "PC-01" in html and "Derecho" in html


@con_db
def test_pdf_del_reporte_es_un_pdf(datos):
    r = datos.get("/reporte.pdf?periodo=rango&desde=2026-09-01&hasta=2026-09-30")
    assert r.status_code == 200
    assert r.content_type == "application/pdf"
    assert r.data[:4] == b"%PDF"
    assert "reporte" in r.headers["Content-Disposition"]


@con_db
def test_excel_del_reporte_trae_una_hoja_por_tabla(datos):
    import openpyxl
    r = datos.get("/reporte.xlsx?periodo=rango&desde=2026-09-01&hasta=2026-09-30")
    assert r.status_code == 200
    libro = openpyxl.load_workbook(io.BytesIO(r.data))
    assert libro.sheetnames == ["Resumen", "Equipos", "Carreras", "Horas", "Dias"]
    equipos = list(libro["Equipos"].iter_rows(values_only=True))
    assert equipos[0] == ("Equipo", "Entradas")
    assert equipos[1] == ("PC-01", 3)


@con_db
def test_la_tarjeta_este_semestre_usa_el_semestre_calendario(datos):
    d = datos.get("/api/stats").get_json()
    # hoy es 2026-09-18 en el entorno de pruebas; el semestre empieza el 1 de agosto,
    # asi que el LOGIN del 20 de agosto cuenta y no hay nada anterior.
    assert d["logins_semestre"] == 5
