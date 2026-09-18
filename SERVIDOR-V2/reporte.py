"""
Reporte de estadisticas de uso del laboratorio para un periodo.

- rango_periodo: traduce "este mes", "mes anterior", "semestre" o un rango
  explicito a fechas y una etiqueta legible.
- calcular_reporte: consulta la bitacora en Postgres y devuelve un dict con
  cifras y tablas. Las marcas de tiempo se guardan en UTC y se convierten a
  la hora de Chihuahua antes de agrupar.
- reporte_xlsx: el mismo contenido como libro de Excel.
"""

import io
from datetime import date, datetime, timedelta

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

ZONA = "America/Chihuahua"
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio",
         "agosto", "septiembre", "octubre", "noviembre", "diciembre"]
PERIODOS = ("mes", "mes-anterior", "semestre", "rango")


class PeriodoInvalido(ValueError):
    """El periodo pedido no se puede interpretar."""


def _mes(fecha):
    return f"{MESES[fecha.month - 1].capitalize()} de {fecha.year}"


def _fecha(texto, nombre):
    try:
        return date.fromisoformat(texto or "")
    except ValueError:
        raise PeriodoInvalido(f"La fecha '{nombre}' debe tener la forma AAAA-MM-DD")


def inicio_semestre(hoy):
    """Enero-julio o agosto-diciembre, el calendario de la UACH."""
    return date(hoy.year, 8, 1) if hoy.month >= 8 else date(hoy.year, 1, 1)


def rango_periodo(periodo, hoy=None, desde=None, hasta=None):
    """Devuelve (desde, hasta, etiqueta) para el periodo pedido."""
    hoy = hoy or date.today()
    if periodo == "mes":
        return date(hoy.year, hoy.month, 1), hoy, _mes(hoy)
    if periodo == "mes-anterior":
        ultimo = date(hoy.year, hoy.month, 1) - timedelta(days=1)
        return date(ultimo.year, ultimo.month, 1), ultimo, _mes(ultimo)
    if periodo == "semestre":
        inicio = inicio_semestre(hoy)
        nombre = "agosto-diciembre" if inicio.month == 8 else "enero-julio"
        return inicio, hoy, f"Semestre {nombre} {hoy.year}"
    if periodo == "rango":
        d, h = _fecha(desde, "desde"), _fecha(hasta, "hasta")
        if d > h:
            raise PeriodoInvalido("La fecha inicial no puede ser posterior a la final")
        return d, h, f"Del {d:%d/%m/%Y} al {h:%d/%m/%Y}"
    raise PeriodoInvalido(f"Periodo desconocido: {periodo}. Usa mes, mes-anterior, semestre o rango")


def _consultar(cur, sql, params):
    cur.execute(sql, params)
    return cur.fetchall()


def calcular_reporte(conn, desde, hasta, etiqueta):
    """Cifras y tablas del periodo [desde, hasta] (fechas locales, inclusive)."""
    # Se compara con la hora local: la marca UTC convertida a Chihuahua.
    filtro = (
        "FROM bitacora_uso b LEFT JOIN alumnos a ON b.matricula = a.cardnumber "
        f"WHERE b.evento = 'LOGIN' AND (b.timestamp AT TIME ZONE 'UTC' AT TIME ZONE '{ZONA}')::date BETWEEN %s AND %s"
    )
    params = (desde, hasta)
    with conn.cursor() as cur:
        total, distintos, dias_act = _consultar(cur,
            f"SELECT COUNT(*), COUNT(DISTINCT b.matricula), COUNT(DISTINCT (b.timestamp AT TIME ZONE 'UTC' AT TIME ZONE '{ZONA}')::date) {filtro}",
            params)[0]
        equipos = _consultar(cur, f"SELECT b.computer_id, COUNT(*) AS n {filtro} GROUP BY b.computer_id ORDER BY n DESC, b.computer_id LIMIT 15", params)
        carreras = _consultar(cur, f"SELECT COALESCE(a.sort1, 'Sin registro'), COUNT(*) AS n {filtro} GROUP BY 1 ORDER BY n DESC, 1 LIMIT 15", params)
        horas = _consultar(cur, f"SELECT EXTRACT(HOUR FROM b.timestamp AT TIME ZONE 'UTC' AT TIME ZONE '{ZONA}')::int AS h, COUNT(*) {filtro} GROUP BY h", params)
        dias = _consultar(cur, f"SELECT (b.timestamp AT TIME ZONE 'UTC' AT TIME ZONE '{ZONA}')::date AS d, COUNT(*) {filtro} GROUP BY d ORDER BY d", params)

    por_hora = [0] * 24
    for h, n in horas:
        por_hora[h] = n
    return {
        "periodo": {"desde": desde.isoformat(), "hasta": hasta.isoformat(), "etiqueta": etiqueta},
        "generado": datetime.now().strftime("%d/%m/%Y %H:%M"),
        "total_entradas": total,
        "alumnos_distintos": distintos,
        "dias_con_actividad": dias_act,
        "promedio_por_dia": round(total / dias_act, 1) if dias_act else 0,
        "equipos": [{"pc": pc, "total": n} for pc, n in equipos],
        "carreras": [{"carrera": c, "total": n} for c, n in carreras],
        "horas": por_hora,
        "dias": [{"fecha": d.isoformat(), "total": n} for d, n in dias],
    }


def horas_pico(datos, cuantas=3):
    """Las horas del dia con mas entradas, como [(hora, total)]."""
    pares = [(h, n) for h, n in enumerate(datos["horas"]) if n]
    return sorted(pares, key=lambda p: (-p[1], p[0]))[:cuantas]


def reporte_xlsx(datos):
    """El reporte como libro de Excel en bytes: una hoja por tabla."""
    libro = openpyxl.Workbook()
    negrita = Font(bold=True)

    def hoja(nombre, encabezados, filas):
        ws = libro.create_sheet(nombre) if libro.sheetnames != ["Sheet"] else libro.active
        ws.title = nombre
        ws.append(encabezados)
        for celda in ws[1]:
            celda.font = negrita
        for fila in filas:
            ws.append(fila)
        for i, enc in enumerate(encabezados, start=1):
            ancho = max([len(str(enc))] + [len(str(f[i - 1])) for f in filas]) + 2
            ws.column_dimensions[get_column_letter(i)].width = min(ancho, 60)
        return ws

    p = datos["periodo"]
    hoja("Resumen", ["Concepto", "Valor"], [
        ("Periodo", p["etiqueta"]),
        ("Desde", p["desde"]),
        ("Hasta", p["hasta"]),
        ("Generado", datos["generado"]),
        ("Entradas de alumnos", datos["total_entradas"]),
        ("Alumnos distintos", datos["alumnos_distintos"]),
        ("Dias con actividad", datos["dias_con_actividad"]),
        ("Promedio de entradas por dia con actividad", datos["promedio_por_dia"]),
    ])
    hoja("Equipos", ["Equipo", "Entradas"], [(e["pc"], e["total"]) for e in datos["equipos"]])
    hoja("Carreras", ["Carrera", "Entradas"], [(c["carrera"], c["total"]) for c in datos["carreras"]])
    hoja("Horas", ["Hora", "Entradas"], [(f"{h:02d}:00", n) for h, n in enumerate(datos["horas"])])
    hoja("Dias", ["Fecha", "Entradas"], [(d["fecha"], d["total"]) for d in datos["dias"]])

    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()
